#!/usr/bin/env python3
"""ساخت نمونه‌های ۳-ردیفی از خروجی‌های واقعی برای تست ایمپورت (گام ۱-ب)"""
import openpyxl, sys

SLICES = [
    ('/tmp/export-melk-tabriz-listings.xlsx', '/tmp/sample-melk-tabriz-listings.xlsx'),
    ('/tmp/export-melk-tabriz-requests.xlsx', '/tmp/sample-melk-tabriz-requests.xlsx'),
    ('/tmp/export-melketun-listings.xlsx', '/tmp/sample-melketun-listings.xlsx'),
    ('/tmp/export-melketun-requests.xlsx', '/tmp/sample-melketun-requests.xlsx'),
]
for src, dst in SLICES:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    # keep header + first 3 data rows: delete rows 5..max
    if ws.max_row > 4:
        ws.delete_rows(5, ws.max_row - 4)
    wb.save(dst)
    print(f"{dst}: {ws.max_row} rows (header+{ws.max_row-1})")
print("done")
